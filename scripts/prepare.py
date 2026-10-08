"""Extract published numbers and calculate 2019-based budget diagnostics.

Run from the repository root: pixi run python -m scripts.prepare
Inputs and their definitions are documented in data/README.md.
"""

import calendar
import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import tomllib
import xarray as xr

from src.budget import anomaly, oh_forcing_from_removal, required_forcing


def calibrated(y, pixels, values):
    """Linear axis calibration; pixel y increases downward."""
    slope = (values[1] - values[0]) / (pixels[1] - pixels[0])
    return values[0] + slope * (np.asarray(y) - pixels[0])


def main():
    config = tomllib.loads(Path("config.toml").read_text())
    raw = Path(config["paths"]["raw"])
    inputs = Path(config["paths"]["transcribed"])
    out = Path(config["paths"]["processed"])
    out.mkdir(parents=True, exist_ok=True)
    years = np.arange(
        config["analysis"]["baseline_year"], config["analysis"]["last_year"] + 1
    )

    # NOAA annual means and Jan-1 to Jan-1 growth are different observables.
    mean = pd.read_csv(
        raw / "noaa/ch4_annmean_gl.txt",
        sep=r"\s+",
        comment="#",
        names=["year", "mean", "mean_unc"],
    )
    growth = pd.read_csv(
        raw / "noaa/ch4_gr_gl.txt",
        sep=r"\s+",
        comment="#",
        names=["year", "growth", "growth_unc"],
    )
    noaa = mean.merge(growth, on="year").set_index("year").loc[years].to_xarray()
    noaa["mass"] = 2.77 * noaa["mean"]
    noaa["growth_mass"] = 2.77 * noaa.growth
    noaa["concentration_anomaly_percent"] = (
        100 * anomaly(noaa["mean"]) / noaa["mean"].sel(year=2019)
    )
    noaa["growth_anomaly"] = anomaly(noaa.growth_mass)
    tau = xr.DataArray(
        [9.0, 10.0, 11.0, 12.0],
        dims="chemical_lifetime",
        coords={"chemical_lifetime": [9.0, 10.0, 11.0, 12.0]},
        name="chemical_lifetime",
    )
    noaa["burden_adjustment"] = anomaly(noaa.mass) / tau
    noaa["required_forcing"] = required_forcing(noaa.growth_mass, noaa.mass, 1 / tau)
    noaa.attrs.update(
        source="NOAA GML files retrieved 2026-10-07/08 UTC; see downloads.jsonl",
        baseline_year=2019,
        conversion_Tg_per_ppb=2.77,
        caveat="Annual surface-mean burden proxy; not native model burden",
    )
    for name in ["mass"]:
        noaa[name].attrs["units"] = "Tg CH4"
    for name in [
        "growth_mass",
        "growth_anomaly",
        "burden_adjustment",
        "required_forcing",
    ]:
        noaa[name].attrs["units"] = "Tg CH4 yr-1"
    noaa.to_netcdf(out / "noaa.nc")

    digit = json.loads((inputs / "he_digitization.json").read_text())
    he = xr.Dataset(coords={"year": digit["year"]})
    for name, spec in digit["figure2"]["series"].items():
        he[name] = (
            "year",
            calibrated(spec["y"], spec["calibration_y"], spec["calibration_value"]),
        )
        he[name].attrs["units"] = spec["units"]
    b = digit["figure2"]["panel_b"]
    gray, pink, red = [
        calibrated(b[name], b["calibration_y"], b["calibration_value"])
        for name in ["gray_top_y", "pink_top_y", "red_top_y"]
    ]
    he["relaxation"] = ("year", gray)
    he["oh_forcing"] = ("year", pink - gray)
    he["emission_forcing_bar"] = ("year", red - pink)
    he["growth"] = ("year", red)
    he["relaxation_anomaly"] = anomaly(he.relaxation)
    he["growth_anomaly"] = anomaly(he.growth)
    he["emission_forcing"] = anomaly(he.emissions)
    a = digit["figure1"]
    for name, y in a["series_y"].items():
        he[name + "_percent"] = (
            "year",
            calibrated(y, a["calibration_y"], a["calibration_value"]),
        )
    he["constant_2021_percent"] = (
        "year",
        np.r_[
            np.nan,
            np.nan,
            calibrated(
                a["constant_2021_y"], a["calibration_y"], a["calibration_value"]
            ),
        ],
    )
    he.attrs.update(
        source=digit["source"],
        method=digit["method"],
        coordinate_error_pixels=2,
        baseline_year=2019,
        caveat="Digitized central estimates; not original posterior data",
    )
    he.to_netcdf(out / "he_digitized.nc")
    he.to_dataframe().to_csv(out / "he_digitized.csv")

    tab = pd.read_csv(inputs / "published_annual.csv")
    annual = (
        tab.drop(columns="source")
        .set_index(["study", "year"])
        .to_xarray()
        .reindex(year=years)
    )
    annual["emission_forcing"] = anomaly(annual.emissions)
    annual["oh_sink_saving"] = -anomaly(annual.oh_removal)
    annual["other_sink_saving"] = -anomaly(annual.other_removal)
    annual["growth"] = annual.emissions - annual.oh_removal - annual.other_removal
    annual["growth_anomaly"] = anomaly(annual.growth)
    annual["oh_forcing_common_burden"] = oh_forcing_from_removal(
        annual.oh_removal, noaa.mass
    )
    annual["oh_forcing_native"] = oh_forcing_from_removal(
        annual.oh_removal, annual.mass
    )
    annual["oh_percent_vs_2019"] = 100 * (
        (1 + annual.oh_percent_vs_prior / 100)
        / (1 + annual.oh_percent_vs_prior.sel(year=2019) / 100)
        - 1
    )
    annual.attrs.update(
        baseline_year=2019,
        caveat="Native OH forcing available for Qu2022; common-burden values are conditional; missing values remain NaN",
    )
    for name in annual.data_vars:
        annual[name].attrs["units"] = (
            "percent"
            if "percent" in name
            else "Tg CH4"
            if name == "mass"
            else "Tg CH4 yr-1"
        )
    annual.to_netcdf(out / "published_annual.nc")
    annual.to_dataframe().to_csv(out / "published_annual.csv")

    # Ciais provides signed removal changes, but no absolute 2019 OH removal.
    digit_c = json.loads((inputs / "ciais_digitization.json").read_text())
    ciais = xr.Dataset(coords={"year": digit_c["year"], "OH_field": ["OH1", "OH2"]})
    spec = digit_c["figure2"]
    emission = calibrated(spec["y"], spec["calibration_y"], spec["calibration_value"])
    ciais["emission_forcing"] = ("year", np.r_[0, emission])
    upper = calibrated(
        spec["sd_upper_cap_y"], spec["calibration_y"], spec["calibration_value"]
    )
    lower = calibrated(
        spec["sd_lower_cap_y"], spec["calibration_y"], spec["calibration_value"]
    )
    ciais["emission_ensemble_sd"] = ("year", np.r_[0, (upper - lower) / 2])
    spec_s = digit_c["figure_s3"]
    signed = xr.DataArray(
        [
            calibrated(
                spec_s[field + "_y"],
                spec_s["calibration_y"],
                spec_s["calibration_value"],
            )
            for field in ciais.OH_field.values
        ],
        dims=["OH_field", "year"],
        coords={"OH_field": ciais.OH_field, "year": ciais.year},
    )
    ciais["oh_sink_saving"] = anomaly(signed)
    tau_oh = xr.DataArray(
        config["analysis"]["oh_lifetime_sensitivity"],
        dims="OH_lifetime",
        coords={"OH_lifetime": config["analysis"]["oh_lifetime_sensitivity"]},
    )
    assumed_oh_removal = noaa.mass.sel(year=2019) / tau_oh - ciais.oh_sink_saving
    ciais["oh_forcing_common_burden"] = oh_forcing_from_removal(
        assumed_oh_removal, noaa.mass
    )
    ciais.attrs.update(
        source=digit_c["source"],
        method=digit_c["method"],
        baseline_year=2019,
        caveat="OH forcing uses NOAA burden and an assumed 2019 OH lifetime; not native posterior forcing",
    )
    for name in ciais.data_vars:
        ciais[name].attrs["units"] = "Tg CH4 yr-1"
    ciais.to_netcdf(out / "ciais.nc")
    ciais.to_dataframe().to_csv(out / "ciais.csv")

    reduction = xr.DataArray(
        config["analysis"]["oh_reduction_scenarios"],
        dims="OH_reduction",
        coords={"OH_reduction": config["analysis"]["oh_reduction_scenarios"]},
    )
    pulse = xr.DataArray(
        (years == 2020).astype(float), dims="year", coords={"year": years}
    )
    laughner = xr.Dataset(
        {"oh_forcing_scenario": noaa.mass.sel(year=2019) * reduction / tau_oh * pulse}
    )
    laughner.attrs.update(
        baseline_year=2019,
        caveat="Uniform prescribed 2020 OH reduction; conditional comparison, not Laughner native output",
    )
    laughner.oh_forcing_scenario.attrs["units"] = "Tg CH4 yr-1"
    laughner.to_netcdf(out / "laughner_scenario.nc")

    ensemble = pd.read_csv(inputs / "qu2022_sensitivity.csv")
    ensemble["oh_forcing_reported"] = ensemble.total_forcing - ensemble.emission_change
    ensemble["emission_fraction"] = ensemble.emission_change / ensemble.total_forcing
    ensemble.to_csv(out / "qu2022_sensitivity.csv", index=False)
    pd.read_csv(inputs / "qu2022_regional.csv").to_csv(
        out / "qu2022_regional.csv", index=False
    )

    # Retain the publisher's exact source numbers, including differences from text.
    book = openpyxl.load_workbook(raw / "peng/figure4.xlsx", data_only=True)
    sheet = book["Fig4a"]
    labels = {
        5: "inversion_net_flux",
        10: "ORCHIDEE",
        11: "LPJ_wsl",
        16: "fire",
        21: "anthropogenic",
        26: "OH_removal",
        31: "growth",
    }
    rows = [
        [label] + [sheet.cell(row, col).value for col in range(1, 7)]
        for row, label in labels.items()
    ]
    pd.DataFrame(
        rows,
        columns=["component", "global", "sd", "30N_90N", "0_30N", "30S_0", "90S_30S"],
    ).to_csv(out / "peng_source_data.csv", index=False)

    grid = xr.open_dataset(raw / "peng/inversion.nc")
    time_year = np.floor(grid.time).astype(int)
    month = np.rint((grid.time - time_year) * 12 + 0.5).astype(int)
    days = xr.DataArray(
        [calendar.monthrange(int(y), int(m))[1] for y, m in zip(time_year, month)],
        dims="time",
        coords={"time": grid.time},
    )
    weighted = grid.fch4_posterior * days
    weighted = weighted.assign_coords(year=("time", time_year.values))
    weights = days.assign_coords(year=("time", time_year.values))
    annual_rate = weighted.groupby("year").sum() / weights.groupby("year").sum()
    annual_rate.attrs.update(
        units="kg CH4 m-2 s-1", description="Day-weighted annual mean net surface flux"
    )
    spatial_delta = anomaly(annual_rate) * 86400 * 1e6  # kg to mg, seconds to days
    spatial_delta.attrs.update(units="mg CH4 m-2 d-1", baseline_year=2019)
    output = xr.Dataset(
        {
            "net_flux_anomaly": spatial_delta,
            "annual_net_flux": (annual_rate * grid.area).sum(["lat", "lon"])
            * 365
            * 86400
            / 1e9,
        }
    )
    output["calendar_year_total"] = (
        (weighted * grid.area).sum(["lat", "lon"]).groupby("year").sum() * 86400 / 1e9
    )
    output.calendar_year_total.attrs["units"] = "Tg CH4 per calendar year"
    output.annual_net_flux.attrs["units"] = "Tg CH4 per 365-day year"
    output.attrs["caveat"] = (
        "Net surface flux, including soil uptake; annual rates use a fixed 365-day year"
    )
    output.to_netcdf(out / "peng_inversion.nc")

    workbook = openpyxl.load_workbook(
        raw
        / "papers/stevensonetal2022_supplement/Stevenson_et_al_2022_Supplementary_Material.xlsx",
        data_only=True,
    )
    sheet = workbook["Table S13ab"]
    records = []
    for row in sheet.iter_rows():
        if row[0].value in [
            "Europe",
            "N America",
            "S Asia",
            "E Asia",
            "ROW",
            "Aviation (extra)",
            "Total",
        ]:
            records.append(
                [row[0].value] + [row[i - 1].value for i in [2, 4, 5, 7, 8, 10, 11, 13]]
            )
    # First seven rows belong to Table S13a (four-model mean).
    pd.DataFrame(
        records[:7],
        columns=[
            "region",
            "NOx",
            "NOx_sd",
            "CO",
            "CO_sd",
            "NMVOC",
            "NMVOC_sd",
            "total",
            "total_sd",
        ],
    ).to_csv(out / "stevenson_4mm.csv", index=False)

    audit = {
        "he_panel_a_vs_b_max_emission_difference_Tg": float(
            abs(he.emission_forcing - he.emission_forcing_bar).max()
        ),
        "peng_day_weighted_365_day_anomaly_Tg": float(
            anomaly(output.annual_net_flux).sel(year=2020)
        ),
        "peng_publisher_anomaly_Tg": float(book["Fig4a"]["A5"].value),
        "peng_calendar_year_anomaly_Tg": float(
            anomaly(output.calendar_year_total).sel(year=2020)
        ),
        "noaa_growth_anomaly_Tg": noaa.growth_anomaly.values.tolist(),
        "burden_adjustment_10yr_Tg": noaa.burden_adjustment.sel(
            chemical_lifetime=10
        ).values.tolist(),
        "qu2022_native_OH_forcing_Tg": float(
            annual.oh_forcing_native.sel(study="Qu2022", year=2020)
        ),
        "qu2022_ensemble_emission_fraction_mean": float(
            ensemble.emission_fraction.mean()
        ),
        "qu2022_ensemble_emission_fraction_population_sd": float(
            ensemble.emission_fraction.std(ddof=0)
        ),
        "ciais_2020_digitization_minus_reported_emissions_Tg": float(
            emission[0] - spec["reported_2020_mean"]
        ),
        "ciais_emission_forcing_Tg": ciais.emission_forcing.values.tolist(),
        "ciais_signed_OH_sink_anomalies_Tg": ciais.oh_sink_saving.values.tolist(),
        "laughner_2020_scenario_3percent_11p1yr_Tg": float(
            laughner.oh_forcing_scenario.sel(
                year=2020, OH_reduction=0.03, OH_lifetime=11.1
            )
        ),
    }
    (out / "audit.json").write_text(json.dumps(audit, indent=2) + "\n")
    report = Path(config["paths"]["report"])
    report.mkdir(exist_ok=True)
    lines = [
        r"\begin{tabular}{lrrr}",
        r"\toprule Study / quantity & 2020 & 2021 & 2022 \\",
        r"\midrule",
    ]
    table_series = [
        ("He emissions (digitized)", he.emission_forcing),
        ("He OH forcing (digitized)", he.oh_forcing),
        ("He relaxation anomaly (digitized)", he.relaxation_anomaly),
        ("Qu 2022 emissions", annual.emission_forcing.sel(study="Qu2022")),
        (
            "Qu 2022 OH forcing, native burden",
            annual.oh_forcing_native.sel(study="Qu2022"),
        ),
        ("Ciais emissions (digitized)", ciais.emission_forcing),
        (
            "Ciais OH-1, conditional",
            ciais.oh_forcing_common_burden.sel(OH_field="OH1", OH_lifetime=11.1),
        ),
        (
            "Ciais OH-2, conditional",
            ciais.oh_forcing_common_burden.sel(OH_field="OH2", OH_lifetime=11.1),
        ),
        ("Qu 2024 emissions", annual.emission_forcing.sel(study="Qu2024")),
        ("Qu 2024 OH removal saving", annual.oh_sink_saving.sel(study="Qu2024")),
        (
            "Qu 2024 OH forcing, common burden",
            annual.oh_forcing_common_burden.sel(study="Qu2024"),
        ),
        (
            "Feng emissions, joint OH",
            annual.emission_forcing.sel(study="Feng2023_joint"),
        ),
        (
            "Feng emissions, fixed OH",
            annual.emission_forcing.sel(study="Feng2023_fixed"),
        ),
        ("NOAA growth anomaly", noaa.growth_anomaly),
        (
            r"Burden adjustment, $\tau_0=10$ yr",
            noaa.burden_adjustment.sel(chemical_lifetime=10),
        ),
    ]
    for label, series in table_series:
        values = [float(series.sel(year=int(y))) for y in years[1:]]
        lines.append(
            label
            + " & "
            + " & ".join(f"{v:.1f}" if np.isfinite(v) else "---" for v in values)
            + r" \\"
        )
    lines += [r"\bottomrule", r"\end{tabular}"]
    (report / "annual_table.tex").write_text("\n".join(lines) + "\n")
    print(json.dumps(audit, indent=2))


if __name__ == "__main__":
    main()
